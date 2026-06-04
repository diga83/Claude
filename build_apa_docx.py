#!/usr/bin/env python3
"""Build an APA 7th-edition formatted .docx from the 6-7 narrative review markdown."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/home/user/Claude/phatic-only-slang-6-7-review.md"
OUT = "/home/user/Claude/Phatic-Only_Slang_6-7_APA7.docx"

TITLE = ('Phatic-Only Slang and the Limits of Meaning-Based Lexical Theory: '
         'A Sociolinguistic Review of the “6-7” Meme')

# ---------- helpers ----------------------------------------------------------
INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|https?://[^\s)]+)')

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0000FF'); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_run(run):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)


def add_md(paragraph, text):
    """Render markdown inline (**bold**, *italic*, URLs) into runs."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()]); style_run(r)
        tok = m.group(0)
        if tok.startswith('**'):
            r = paragraph.add_run(tok[2:-2]); r.bold = True; style_run(r)
        elif tok.startswith('http'):
            add_hyperlink(paragraph, tok, tok)
        elif tok.startswith('*'):
            r = paragraph.add_run(tok[1:-1]); r.italic = True; style_run(r)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:]); style_run(r)


def fmt(p, *, indent=False, hanging=False, align=None, space_after=0):
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Inches(0.5)
    if hanging:
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)
    return p

# ---------- document setup ---------------------------------------------------
doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 2.0
normal.paragraph_format.space_after = Pt(0)

sec = doc.sections[0]
for side in ('top', 'bottom', 'left', 'right'):
    setattr(sec, f'{side}_margin', Inches(1))

# page number in header, right aligned (APA student paper)
hdr = sec.header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = hp.add_run()
style_run(run)
fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

# ---------- title page -------------------------------------------------------
for _ in range(3):
    fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run(TITLE); r.bold = True; style_run(r)
fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)  # blank line
for line in ['[Author Name]',
             '[Department], [Institution]',
             '[Course Number: Course Title]',
             '[Instructor Name]',
             'June 4, 2026']:
    p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(line); style_run(r)

# ---------- parse markdown ---------------------------------------------------
with open(SRC, encoding='utf-8') as f:
    lines = f.read().split('\n')

# cut body at References (we supply an APA reference list instead)
end = next(i for i, l in enumerate(lines) if l.strip() == '## References')
body = lines[:end]

def page_break():
    doc.add_page_break()

abstract_done = False
body_started = False
i = 0
while i < len(body):
    line = body[i].rstrip()
    i += 1
    if not line or line == '---':
        continue
    if line.startswith('# '):          # main title -> skip (on title page)
        continue
    if line.startswith('## Abstract'):
        page_break()
        p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run('Abstract'); r.bold = True; style_run(r)
        # abstract body (next non-empty line)
        while i < len(body) and not body[i].strip():
            i += 1
        if i < len(body):
            p = fmt(doc.add_paragraph())
            add_md(p, body[i].strip())
            i += 1
        # keywords
        p = fmt(doc.add_paragraph(), indent=True)
        rk = p.add_run('Keywords: '); rk.italic = True; style_run(rk)
        rk2 = p.add_run('phatic communion, phatic function, lexical semantics, slang, '
                        'indexicality, internet linguistics, brainrot, Generation Alpha')
        style_run(rk2)
        abstract_done = True
        continue
    if line.startswith('## '):         # Level 1 heading
        if not body_started:
            page_break()
            p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(TITLE); r.bold = True; style_run(r)
            body_started = True
        text = line[3:].strip()
        p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(text); r.bold = True; style_run(r)
        continue
    if line.startswith('### '):        # Level 2 heading
        text = line[4:].strip()
        p = fmt(doc.add_paragraph())
        r = p.add_run(text); r.bold = True; style_run(r)
        continue
    if line.startswith('- '):          # bullet item
        p = fmt(doc.add_paragraph(), hanging=True)
        rb = p.add_run('•  '); style_run(rb)
        add_md(p, line[2:].strip())
        continue
    if line.startswith('*') and line.endswith('*') and not line.startswith('**'):
        p = fmt(doc.add_paragraph())
        add_md(p, line)
        continue
    # normal body paragraph
    p = fmt(doc.add_paragraph(), indent=True)
    add_md(p, line)

# ---------- APA 7 references -------------------------------------------------
page_break()
p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run('References'); r.bold = True; style_run(r)

REFS = [
 'Axios. (2025, October 6). *TikTok’s “6-7” and Gen Alpha classroom bans*. https://www.axios.com/2025/10/06/tiktok-6-7-gen-alpha-classroom-bans',
 'Blakemore, D. (2002). *Relevance and linguistic meaning: The semantics and pragmatics of discourse markers*. Cambridge University Press.',
 'Davidson, D. (1967). Truth and meaning. *Synthese, 17*(3), 304–323. https://doi.org/10.1007/BF00485035',
 'Dawkins, R. (1976). *The selfish gene*. Oxford University Press.',
 'Dictionary.com. (2025, October 28). *Dictionary.com’s 2025 word of the year is “67.”* https://www.dictionary.com/articles/word-of-the-year-2025',
 'Eble, C. (1996). *Slang and sociability: In-group language among college students*. University of North Carolina Press.',
 'Eble, C. (2004). Slang [Course reading]. University of Pennsylvania. https://www.ling.upenn.edu/~beatrice/110/readings/eble-2004.pdf',
 'Eckert, P. (2008). Variation and the indexical field. *Journal of Sociolinguistics, 12*(4), 453–476. https://doi.org/10.1111/j.1467-9841.2008.00374.x',
 'Eckert, P. (2018). *Meaning and linguistic variation: The third wave in sociolinguistics*. Cambridge University Press.',
 'Firth, J. R. (1957). A synopsis of linguistic theory, 1930–1955. In *Studies in linguistic analysis* (pp. 1–32). Blackwell.',
 'Frege, G. (1948). Sense and reference (M. Black, Trans.). *The Philosophical Review, 57*(3), 209–230. https://doi.org/10.2307/2181485 (Original work published 1892)',
 'Fridland, V. (2026, April 15). Slanguage: How “6-7” makes sense even though it means nothing. *The Conversation*. https://theconversation.com/slanguage-how-6-7-makes-sense-even-though-it-means-nothing-270006',
 'Georgetown University. (2025). *Linguistics professor explains the “six-seven” meme*. https://www.georgetown.edu/news/six-seven-meme-linguistics/',
 'Goldberg, A. E. (1995). *Constructions: A construction grammar approach to argument structure*. University of Chicago Press.',
 'Gutzmann, D. (2015). *Use-conditional meaning: Studies in multidimensional semantics*. Oxford University Press.',
 'Halliday, M. A. K. (1976). Anti-languages. *American Anthropologist, 78*(3), 570–584. https://doi.org/10.1525/aa.1976.78.3.02a00050',
 'IBTimes UK. (2025, November 25). *Keir Starmer apologises after being told off by teacher for doing “6-7” meme during school visit*. https://www.ibtimes.co.uk/keir-starmer-apologises-after-being-told-off-teacher-doing-6-7-meme-during-school-visit-1757929',
 'Jakobson, R. (1960). Linguistics and poetics. In T. A. Sebeok (Ed.), *Style in language* (pp. 350–377). MIT Press.',
 'Jurafsky, D., & Martin, J. H. (2023). Lexical semantics. In *Speech and language processing* (3rd ed. draft, Chapter 19). https://web.stanford.edu/class/linguist1/Rdgs/JM19.pdf',
 'Kaplan, D. (1999). *The meaning of “ouch” and “oops”: Explorations in the theory of meaning as use* [Unpublished manuscript]. University of California, Los Angeles.',
 'Malinowski, B. (1923). The problem of meaning in primitive languages. In C. K. Ogden & I. A. Richards, *The meaning of meaning* (pp. 296–336). Kegan Paul.',
 'McCulloch, G. (2019). *Because internet: Understanding the new rules of language*. Riverhead Books.',
 'Merriam-Webster. (n.d.). Six seven. In *Merriam-Webster slang dictionary*. Retrieved June 4, 2026, from https://www.merriam-webster.com/slang/six-seven',
 'Miller, V. (2008). New media, networking and phatic culture. *Convergence, 14*(4), 387–400. https://doi.org/10.1177/1354856508094659',
 'NPR. (2025, October 31). *What does “67” mean? Dictionary.com’s 2025 word of the year has no definition*. https://www.npr.org/2025/10/31/nx-s1-5593052/67-dictionary-word-of-the-year-2025',
 'Oxford University Press. (2023). *“Rizz” crowned Oxford word of the year 2023*. https://corp.oup.com/news/rizz-crowned-oxford-word-of-the-year-2023/',
 'Oxford University Press. (2024). *“Brain rot” named Oxford word of the year 2024*. https://corp.oup.com/news/brain-rot-named-oxford-word-of-the-year-2024/',
 'Potts, C. (2005). *The logic of conventional implicatures*. Oxford University Press.',
 'Potts, C. (2007). The expressive dimension. *Theoretical Linguistics, 33*(2), 165–198. https://doi.org/10.1515/TL.2007.011',
 'Silverstein, M. (2003). Indexical order and the dialectics of sociolinguistic life. *Language & Communication, 23*(3–4), 193–229. https://doi.org/10.1016/S0271-5309(03)00013-2',
 'Skrilla. (2025). *Doot doot (6 7)* [Song]. Priority Records.',
 'Smithsonian Magazine. (2025, October 30). *The term “67” is impossible to define. It just became Dictionary.com’s word of the year for 2025*. https://www.smithsonianmag.com/smart-news/the-term-67-is-impossible-to-define-it-just-became-dictionary-coms-word-of-the-year-for-2025-180987650/',
 'Stanford Encyclopedia of Philosophy. (n.d.). *Word meaning*. https://plato.stanford.edu/entries/word-meaning/',
 'Today. (2024, June 27). *What does “skibidi” mean?* https://www.today.com/parents/teens/skibidi-slang-meaning-rcna153557',
 'University of Colorado Boulder. (2021, May 3). *To “yeet” or not to “yeet”: African American slang beyond Vine*. https://www.colorado.edu/linguistics/2021/05/03/yeet-or-not-yeet-african-american-slang-beyond-vine',
 'WHYY. (2025). *The meaning of “6-7” and its Philadelphia roots*. https://whyy.org/articles/meaning-6-7-skrilla-philadelphia/',
 'Wikipedia contributors. (n.d.-a). 6-7 (meme). In *Wikipedia*. Retrieved June 4, 2026, from https://en.wikipedia.org/wiki/6-7_(meme)',
 'Wikipedia contributors. (n.d.-b). Doot doot (6 7). In *Wikipedia*. Retrieved June 4, 2026, from https://en.wikipedia.org/wiki/Doot_Doot_(6_7)',
 'Wittgenstein, L. (1953). *Philosophical investigations* (G. E. M. Anscombe, Trans.). Blackwell.',
]
for ref in REFS:
    p = fmt(doc.add_paragraph(), hanging=True)
    add_md(p, ref)

# ---------- appendix: methodological note ------------------------------------
page_break()
p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run('Appendix'); r.bold = True; style_run(r)
p = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run('Methodological Note on Sourcing and Verification'); r.bold = True; style_run(r)
NOTE = ('This review was compiled in June 2026. Claims were gathered through a fan-out '
        'web search across primary and secondary sources and then subjected to a dedicated '
        'adversarial verification pass, in which independent counter-searches actively sought '
        'to refute each medium- and low-confidence claim. No claim was struck. The empirical '
        'core—Dictionary.com’s rationale and the Steve Johnson quotation; the Skrilla '
        'release date; the Jones, Fridland, Aleksic, and Gordon commentary; the Merriam-Webster '
        'slang entry; and the Dannenbring, Starmer, and “67 Kid” anecdotes—was '
        'confirmed across multiple independent sources. Quotations from primary theoretical texts '
        '(Malinowski, 1923; Jakobson, 1960; Wittgenstein, 1953, §43; Kaplan, 1999; Potts, '
        '2005, 2007; Firth, 1957) were verified verbatim; the Wittgenstein §43 quotation '
        'follows the Anscombe translation. Two citation corrections were applied during '
        'verification: Firth’s “company it keeps” line was re-sourced to “A '
        'Synopsis of Linguistic Theory, 1930–1955” (in Studies in Linguistic Analysis, '
        '1957) rather than Papers in Linguistics, and “indexical field” was credited to '
        'Eckert (2008) rather than the 2018 book. One residual caveat remains: page-level '
        'verification against print editions of the theoretical primary texts is recommended '
        'before formal publication, and the 10-67 police-code etymology of “6-7” remains '
        'unverified speculation.')
p = fmt(doc.add_paragraph(), indent=True)
add_md(p, NOTE)

doc.save(OUT)
print('Saved', OUT)
