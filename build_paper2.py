# -*- coding: utf-8 -*-
"""Build the clock-it narrative review in a single-column, published-article format
matching the user's preferred reference layout (Times New Roman 11pt, justified,
single-spaced with 6pt after, numbered descriptive headings, structured abstract,
run-in subheadings, embedded figures, a 4-column claims table, and 10pt
hanging-indent references)."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = "/home/user/Claude/Clock-It_Narrative-Review.docx"
FIG = "/home/user/Claude/figs"

doc = Document()

# ---- base / Normal ----
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
rpr.set(qn("w:eastAsia"), "Times New Roman")
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
pf.space_after = Pt(6)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for s in doc.sections:
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)

J = WD_ALIGN_PARAGRAPH.JUSTIFY
C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT

def _fmt(r, size=11, bold=False, italic=False, color=None):
    r.font.name = "Times New Roman"; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic
    if color: r.font.color.rgb = color

def P(text="", *, align=J, size=11, bold=False, italic=False,
      before=0, after=6, indent=None):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if indent is not None:
        p.paragraph_format.first_line_indent = Inches(indent)
    if text:
        _fmt(p.add_run(text), size=size, bold=bold, italic=italic)
    return p

def RUNIN(label, body, *, size=11):
    """Run-in bold sublabel followed by normal text, justified."""
    p = doc.add_paragraph(); p.alignment = J
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _fmt(p.add_run(label + " "), size=size, bold=True)
    _fmt(p.add_run(body), size=size)
    return p

def MIXED(parts, *, align=J, size=11, after=6, before=0):
    """parts: list of (text, bold, italic)."""
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for t, b, i in parts:
        _fmt(p.add_run(t), size=size, bold=b, italic=i)
    return p

def H(text):
    """Numbered section heading: bold 13pt, left."""
    p = doc.add_paragraph(); p.alignment = L
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _fmt(p.add_run(text), size=13, bold=True)
    return p

def FIGURE(path, caption, width=6.5):
    p = doc.add_paragraph(); p.alignment = C
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph(); cap.alignment = C
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # bold "Figure N." + normal rest
    head, _, rest = caption.partition(". ")
    _fmt(cap.add_run(head + ". "), size=9, bold=True)
    _fmt(cap.add_run(rest), size=9)

def TBL_CAP(text):
    p = doc.add_paragraph(); p.alignment = L
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    head, _, rest = text.partition(". ")
    _fmt(p.add_run(head + ". "), size=10, bold=True)
    _fmt(p.add_run(rest), size=10, italic=False)

def TBL_NOTE(text):
    p = doc.add_paragraph(); p.alignment = L
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(8)
    _fmt(p.add_run("Note. "), size=9, italic=True)
    _fmt(p.add_run(text), size=9)

def cell_text(cell, text, *, bold=False, size=9.5):
    cell.text = ""
    para = cell.paragraphs[0]; para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    para.paragraph_format.space_after = Pt(0)
    _fmt(para.add_run(text), size=size, bold=bold)

# unicode shortcuts
EN = "–"; LQ = "“"; RQ = "”"; AP = "’"

# ============================================================
# TITLE BLOCK
# ============================================================
MIXED([("Clocking the Tea: A Narrative Review of the Verb ", False, False),
       ("to clock", False, True),
       (" as a Site of Semantic Specialization, Enregisterment, and Racialized "
        "Appropriation", False, False)],
      align=C, size=16, after=10)
P("[Author Name]¹*", align=C, size=12, after=2)
P("¹ [Department / Faculty, Institution, City, Country]    * Corresponding author: "
  "[author@institution.edu]", align=C, size=10, after=10)

# ============================================================
# ABSTRACT (structured, run-in bold labels)
# ============================================================
P("Abstract", align=L, size=12, bold=True, after=4)
MIXED([
    ("Background. ", True, False),
    (f"The imperative {LQ}clock it{RQ} and the verb to clock {EN} meaning to perceive, "
     "recognize, or expose a person or a concealed truth {EN} have become salient in "
     "anglophone digital culture, yet the item sits at the intersection of three processes "
     "sociolinguistics usually studies separately: documented semantic change, "
     "enregisterment, and the racialized appropriation of minoritized vernacular. ", False, False),
    ("Objectives. ", True, False),
    ("This review synthesizes the scholarly treatment of the item across historical "
     "lexicography, transgender studies, the sociolinguistics of Black and Latinx queer "
     "vernacular, and critical media studies of drag. ", False, False),
    ("Method. ", True, False),
    ("A narrative (interpretive) review restricted to peer-reviewed scholarship and "
     "scholarly lexicography, with a direct full-text lexical check of seven central sources. ",
     False, False),
    ("Findings. ", True, False),
    ("The perception sense is lexicographically attested from 1929, roughly half a century "
     "before its ballroom association, so the trans and ballroom uses are best analyzed as a "
     "semantic specialization rather than an origin; within transgender studies clocking is "
     "theorized as the antonym of passing, with recent work inverting the usual valuation; "
     "and the item participates in a documented pathway of appropriation and indexical "
     "erasure. A lexical check finds the lemma clock in only one of seven full-text sources. ",
     False, False),
    ("Conclusions. ", True, False),
    (f"Unlike its siblings shade, reading, and tea, the item clock has received no dedicated "
     "study; the review specifies how to close that gap.", False, False),
], align=J, size=10, after=4)
MIXED([("Keywords: ", False, True),
       ("clock; enregisterment; African American Vernacular English; transgender studies; "
        "ballroom culture; semantic change; linguistic appropriation; raciolinguistics; "
        "lexicography; drag", False, False)],
      align=J, size=10, after=10)

# ============================================================
# 1. INTRODUCTION
# ============================================================
H("1. Introduction")
P(f"By the mid-2020s, the imperative {LQ}clock it{RQ} had become a conspicuous feature of "
  "anglophone digital culture, circulating on TikTok, in reality-television commentary, and "
  "across mainstream journalism as a way of flagging that a hidden truth has been perceived "
  "and named. Its trajectory is, on the surface, familiar: a piece of minoritized vernacular "
  "acquires mass currency and, in the process, sheds much of the context that gave it "
  "meaning. Yet the verb to clock is an unusually instructive case, because it sits at the "
  "intersection of three processes that sociolinguistics has tended to study separately: "
  "ordinary semantic change documented in the historical lexicographic record; the "
  "enregisterment of a minoritized variety with a characterological figure (Agha, 2003; "
  "Ilbury, 2020); and the racialized appropriation of African American Vernacular English "
  "(AAVE) and ballroom vocabulary into a putatively neutral "
  f"{LQ}internet language{RQ} (Crowley, 2025; Ilbury & Walcott, 2026).")
P("The verb denotes, in its core modern sense, an act of recognition: to notice, to read, "
  "or to expose. In transgender contexts it carries a more specific and higher-stakes "
  "meaning: to be clocked is to be perceived as transgender, the experiential and political "
  "opposite of passing (Schiffer, 2022; Serano, 2007; Wong, 2025). In Black and Latinx "
  "ballroom and drag culture it names the detection of a concealed truth or flaw and stands "
  "alongside the better-studied practices of reading and throwing shade (Bailey, 2013; "
  "Barrett, 2017).")
P("This article asks three questions: (a) what does the historical and lexicographic record "
  "establish about the origins and development of the perception sense; (b) how has the "
  "transgender and ballroom sense been theorized; and (c) how has the item figured in "
  "scholarship on mediatization and appropriation? The central argument is that to clock "
  "rewards attention precisely because the scholarship bearing on it is fragmented across "
  "sub-disciplines that rarely cite one another. Read together, these literatures yield a "
  "coherent genealogy and expose a clear gap for future empirical work.")

# ============================================================
# 2. METHOD
# ============================================================
H("2. Method")
P("This is a narrative (interpretive) review rather than a systematic review; its aim is "
  "synthesis across heterogeneous literatures rather than exhaustive enumeration "
  "(Greenhalgh et al., 2018). The procedure, summarized in Figure 1, combined a fan-out "
  "search across five thematic angles with adversarial verification and a direct full-text "
  "check of the most load-bearing sources.")
FIGURE(f"{FIG}/fig1.png",
       "Figure 1. The structured narrative-review and verification workflow, from scoping "
       "five search angles through parallel search, source retrieval, full-text "
       "verification, and cited synthesis.")
RUNIN("Source identification.",
      "Sources were located through iterative database and web searching across five angles "
      "{} historical lexicography; transgender studies; ballroom, drag, and AAVE "
      "sociolinguistics; the semantics and regional distribution of the verb; and "
      "media-studies treatments of drag {} using the lemmas clock, clocked, clocking, and "
      f"{LQ}clock it{RQ} in combination with terms such as passing, reading, shade, "
      "enregisterment, and appropriation.".format(EN, EN))
RUNIN("Source selection and exclusion.",
      "To meet the evidentiary standard of an academic review, the base was restricted to "
      "peer-reviewed publications, scholarly monographs and edited volumes from university "
      "and established academic presses, and recognized scholarly lexicography (the Oxford "
      f"English Dictionary; Green{AP}s Dictionary of Slang; the New Partridge Dictionary of "
      "Slang; Berrey & Van den Bark, 1942). General-purpose dictionaries and journalistic "
      "sources were excluded from the evidentiary base; where popular claims are mentioned "
      "they are flagged and treated as hypotheses, not findings. Crowd-sourced "
      "encyclopaedias were excluded entirely.")
RUNIN("Lexical verification.",
      "For a convenience set of seven full-text sources central to the argument {} Wong "
      "(2025), Ilbury (2020), Ilbury and Walcott (2026), Crowley (2025), McKinnon (2017), "
      "Simmons (2014), and Zimman (2020) {} the full published or author-accepted texts were "
      "examined directly and searched for the lemma clock. The result motivates the "
      f"review{AP}s central claim and is reported in Sections 8 and 9.".format() .replace("{} ","{} ").format() if False else
      "For a convenience set of seven full-text sources central to the argument "
      f"{EN} Wong (2025), Ilbury (2020), Ilbury and Walcott (2026), Crowley (2025), "
      f"McKinnon (2017), Simmons (2014), and Zimman (2020) {EN} the full published or "
      "author-accepted texts were examined directly and searched for the lemma clock. The "
      f"result motivates the review{AP}s central claim and is reported in Sections 8 and 9.")
RUNIN("Synthesis.",
      "Findings were organized thematically rather than chronologically, integrating the "
      "lexicographic chronology with the conceptual and political analyses of the cultural "
      "literatures. A confidence-rated summary of principal claims appears in Table 2.")
MIXED([("A limitation must be stated at the outset. ", False, True),
       (f"The dated lexicographic citations in Section 3 derive from scholarly slang "
        "dictionaries consulted through catalogue and secondary evidence and should be "
        "confirmed against the subscription or print originals before being quoted verbatim. "
        "Where the literature supports only an inference rather than a direct statement "
        f"{EN} for example, linking a monograph on surveillance to the slang lexeme clock "
        f"{EN} this is stated explicitly.", False, False)], align=J)

# ============================================================
# 3. LEXICAL SUBSTRATE
# ============================================================
H("3. The Lexical Substrate: Etymology and the Perception Sense")
P("The noun clock entered English in the late fourteenth century, ultimately from Medieval "
  f"Latin clocca {LQ}bell{RQ} (probably of Celtic origin), by way of Middle Dutch and Old "
  "North French (Oxford English Dictionary [OED], n.d.). From this base, slang lexicography "
  "records several distinct verbal developments that popular accounts routinely conflate "
  "(Table 1).")
MIXED([("A first branch derives from the early-twentieth-century use of clock as a slang "
        "noun for the human face, by analogy with the face of a timepiece rather than "
        "through rhyming slang; from this noun comes the verb to clock ", False, False),
       (f"{LQ}to strike or punch,{RQ}", False, False),
       (" recorded from the early 1940s and frequently labeled Australian in origin "
        f"(the American idiom to clean someone{AP}s clock belongs to the same family). "
        "This pugilistic sense is etymologically and semantically distinct from the "
        "perception sense that is the subject of this review.", False, False)], align=J)
MIXED([("The perception sense {} to clock as ".format(EN), False, False),
       (f"{LQ}to catch sight of, to notice, to watch, to recognize{RQ}", False, False),
       (f" {EN} is the pivotal one. The New Partridge Dictionary of Slang dates it to 1929, "
        "and a further early attestation appears in Berrey and Van den Bark (1942), glossed "
        f"as {LQ}to look at; see.{RQ} Green{AP}s Dictionary of Slang offers the fullest "
        f"gloss: {LQ}to see, to recognize, to notice, to watch, to understand, to work "
        f"something out{RQ} (Green, n.d.). The crucial datum is chronological: the general "
        "perception sense is securely attested roughly half a century before the ballroom "
        "and transgender uses with which the item is popularly associated.", False, False)],
      align=J)
P("Two uncertainties remain and should not be smoothed over. The first concerns regional "
  f"origin: although the perception sense is widely felt to be British or Commonwealth "
  f"colloquial (as in {LQ}I clocked him the moment he walked in{RQ}), lexicographic evidence "
  "variously tags it as originally North American, producing a mismatch the accessible "
  "sources do not resolve. The second concerns the earliest citation, with the 1929 "
  "Partridge dating and the 1942 attestation as the two leading candidates. Neither "
  "uncertainty disturbs the central chronological point.")

TBL_CAP("Table 1. Documented senses of the verb to clock and their approximate first attestation.")
rows = [
    ("Sense", "Approx. first attestation", "Source family"),
    ("Noun: timepiece (< L. clocca 'bell')", "c. 1350–1400", "OED"),
    ("Noun: human face (timepiece analogy)", "early 20th c.", "OED; Green's"),
    ("Verb: to strike, punch (< 'face' noun)", "c. 1941 (orig. Australian)", "OED"),
    ("Verb: to notice, see, recognize, watch", "1929; 1942", "New Partridge; Green's"),
    ("Verb: to register / record (a time, a speed)", "20th c.", "OED; Green's"),
    ("Verb (specialized): to recognize someone as transgender",
     "later 20th c. (not lexicographically dated)", "Trans studies; Homosaurus"),
]
t1 = doc.add_table(rows=0, cols=3); t1.style = "Table Grid"; t1.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, r in enumerate(rows):
    cells = t1.add_row().cells
    for c, txt in zip(cells, r):
        cell_text(c, txt, bold=(i == 0), size=9.5)
TBL_NOTE("Datings derive from scholarly slang lexicography and should be confirmed against "
         "the subscription or print originals (see Section 2).")

# ============================================================
# 4. SPECIALIZATION
# ============================================================
H("4. From Perception to Detection: A Semantic Specialization")
P("A recurrent claim in popular writing is that to clock originated in the Black and Latinx "
  "ballroom culture of 1980s New York. The lexicographic evidence renders this untenable as "
  "an account of origin while clarifying what is nonetheless correct in it. The verb of "
  "perception predates the ballroom scene; what ballroom and transgender communities "
  "contributed was not the verb but a high-stakes specialization of it, narrowing a general "
  "sense of noticing toward a specific and consequential referent: the detection that a "
  "person is transgender, in drag, or otherwise not "
  f"{LQ}real{RQ} in the ballroom sense of realness (Bailey, 2013). As Figure 2 shows, this "
  "is a clean instance of narrowing followed by later re-generalization as the item diffused "
  "outward into the generic contemporary sense of naming a concealed truth.")
FIGURE(f"{FIG}/fig2.png",
       "Figure 2. The semantic genealogy of to clock: a separate 'strike' branch from the "
       "'face' noun (left) versus the perception sense (center) that narrows into the "
       "trans/ballroom specialization and then re-generalizes into mainstream 'clock it.'",
       width=5.6)
P("Treating the ballroom usage as origin obscures this trajectory; treating it as "
  "specialization makes the item's history intelligible and reconciles the lexicographic and "
  "cultural literatures, which otherwise appear to contradict one another.")

# ============================================================
# 5. TRANS STUDIES
# ============================================================
H("5. Clocking in Transgender Studies: Recognition, Passing, and Surveillance")
MIXED([("It is in transgender studies that clocking has received its most sustained "
        "conceptual treatment, even though the field generally adopts the term as community "
        f"vernacular rather than as an object of lexical analysis. The pivotal contribution "
        f"is Wong{AP}s (2025) analysis of the ambivalences of trans recognition. Drawing on "
        "Honneth, Wong defines the term directly: ", False, False),
       (f"{LQ}Clocking is literally recognition of a person{AP}s transness, which implies "
        f"that person{AP}s failing to pass, being spotted in their transness, and losing "
        f"realness{RQ}", False, False),
       (" (Wong, 2025), and elaborates that ", False, False),
       (f"{LQ}being clocked designates being spotted as trans, losing realness, failing to "
        f"pass{RQ}", False, False),
       (" while clocking ", False, False),
       (f"{LQ}inscribes the demands of cis-normativity within an act of seeing a person as "
        f"trans{RQ}", False, False),
       (" (Wong, 2025).", False, False)], align=J)
MIXED([("Wong then performs a deliberate inversion of the term's usual valuation: because "
        "the criteria of passing are themselves cis-normative, ", False, False),
       (f"{LQ}passing, not being clocked, is misrecognition; consequently, recognition is "
        f"clocking{RQ}", False, False),
       (" (Wong, 2025), a position condensed into the formula that ", False, False),
       (f"{LQ}recognition is clocking, and clocking is recognition.{RQ}", False, False),
       (" This stands as the sharpest scholarly disagreement in the corpus, directly "
        "contesting the lay framing in which being clocked is uniformly a harm. It is also, "
        "notably, the only source in the assembled set that uses the lexeme clock at all "
        "(Section 8).", False, False)], align=J)
P("Schiffer's (2022) ethnographic study brings the term into interactional sociology. "
  "Reflecting on fieldwork conducted as a transgender woman, Schiffer shows that whether or "
  "not she was clocked materially restructured her interactions with predominantly cisgender "
  "participants, and theorizes the resulting dynamic through Goffman, coining cistress for "
  "the guilt or anger that cisgender interlocutors display when cisnormative assumptions are "
  "disrupted. Serano's (2007) widely cited account {} a trade-press book by a "
  "scholar-activist, flagged as such {} frames being read or clocked as a product of the "
  f"{LQ}cis assumption,{RQ} after which observers engage in {LQ}ungendering.{RQ} The "
  "conceptual antonym is theorized by Beauchamp (2019), whose monograph on transgender "
  "politics and U.S. surveillance analyzes going stealth and argues that heightened "
  "visibility can render transgender people more rather than less surveillable; the link to "
  "the slang lexeme clock is conceptual rather than lexical. Finally, the emerging programme "
  "of trans linguistics (Zimman, 2020) theorizes how transness is perceived through vocal "
  "and discursive practice, yet {} tellingly {} does not foreground the lexeme clock."
  .replace("{} ", f"{EN} ", 4))

# ============================================================
# 6. READING AND SHADE
# ============================================================
H("6. Clocking among Reading and Shade: The Pragmatics of Black Queer Vernacular")
P("The item belongs to a Black and Latinx queer vernacular system whose other members are "
  "comparatively well studied. The foundational linguistic scholarship is Barrett's work on "
  "African American drag performance (Barrett, 1998, 2017), which analyzes the strategic "
  "styleswitching through which Black drag queens index multiple social personae, and which "
  "constitutes the intellectual home for the cluster reading, shade, and clocking. Bailey's "
  "(2013) ethnography establishes the broader frame, treating ballroom vernacular as a "
  "practice of survival, kinship, and resistance and supplying the authoritative account of "
  "realness {} the achievement against which being clocked is the failure, with material "
  "consequences for safety, housing, and employment.".replace("{} ", f"{EN} "))
P("The pragmatic literature distinguishes the cluster with some precision: clocking is the "
  "detection of a concealed truth or flaw; reading is the artful delivery of a pointed "
  "insult; and shade is the indirect, implicational form of reading. Reading has been "
  "analyzed as mock impoliteness that builds in-group solidarity rather than division: "
  "McKinnon (2017) examines reading in the backstage talk of four drag queen performers and "
  f"argues that what appears as insult functions to build a shared {LQ}thick skin,{RQ} while "
  f"McKinnon (2023) and Simmons (2014) extend the analysis to the mediatized {LQ}Reading "
  f"Challenge{RQ} and to the speech codes of RuPaul{AP}s Drag Race respectively. It is "
  "telling that even within this comparatively developed literature, reading and shade are "
  "defined far more sharply than clocking, which is glossed only in passing where it appears "
  "at all.")

# ============================================================
# 7. MEDIATIZATION & APPROPRIATION
# ============================================================
H("7. Enregisterment, Mediatization, and Racialized Appropriation")
P("The contemporary circulation of clock it is best understood through the sociolinguistic "
  "literature on enregisterment and appropriation. Ilbury (2020) shows, in a corpus of "
  "15,804 tweets from ten gay British men, how AAVE features are stylized to project a "
  f"recognizable {LQ}Sassy Queen{RQ} persona that rests on an essentialized image of Black "
  "women, demonstrating the enregisterment of AAVE with a characterological figure rather "
  "than with its actual speaker community. Building on this, Ilbury and Walcott (2026) "
  f"analyze a corpus of 178 TikTok videos in which AAVE features are recontextualized as "
  f"{LQ}internet{RQ} or {LQ}Gen Z{RQ} language; crucially, they show that two competing "
  f"metalinguistic discourses {EN} a concern over the {LQ}indexical erasure{RQ} of AAVE "
  f"versus claims of a new register of {LQ}internet language{RQ} {EN} coexist among users "
  "themselves, so the contradiction is internal to the data. Figure 3 places this diffusion "
  "in historical perspective.")
FIGURE(f"{FIG}/fig3.png",
       "Figure 3. Transmission timeline of the perception sense, from its lexicographic "
       "attestation (1929) through ballroom/trans specialization, mediatization on RuPaul's "
       "Drag Race, and social-media diffusion to mainstream 'clock it.'")
P("Crowley (2025) extends the analysis to the racial politics of attribution within "
  "transgender communities. In ethnographic interviews with seven transgender South "
  "Carolinians, linguistic innovations traced to Black trans communities were characterized "
  f"as {LQ}slang{RQ} (e.g., sis, queen), while those attributed to white communities were "
  f"characterized as {LQ}terminology{RQ} (e.g., demigender, nonbinary), and participants "
  f"described a trajectory by which Black-originated terms travel outward until they {LQ}just "
  f"become queer slang,{RQ} losing their racial attribution. This is precisely the mechanism "
  "an item such as clock it undergoes. Critical media studies situate the diffusion within a "
  "political economy of commodification: Goldmark (2015) examines how the language of "
  "inclusion in RuPaul's Drag Race operates against an unmarked norm of English proficiency, "
  "while Hodes and Sandoval (2018) argue that the franchise commodifies and depoliticizes "
  "drag. These accounts stand in tension with celebratory framings of the same diffusion as "
  "linguistic creativity {} the recurring split between innovation and erasure, subversion "
  "and reinscription, that an adequate account of clock it must hold in view."
  .replace("{} ", f"{EN} "))

# ============================================================
# 8. LACUNA
# ============================================================
H("8. The Lexical Lacuna: Evidence from a Full-Text Check")
P("Despite the rich treatment of its lexical siblings, the item clock has attracted no "
  "dedicated empirical or lexicological study; it is everywhere adjacent and nowhere "
  "central. The point is not merely impressionistic. In the assembled set of seven full-text "
  "sources closest to the topic, the lemma clock appears only in Wong (2025) and is absent "
  "from the remaining six {} Ilbury (2020), Ilbury and Walcott (2026), Crowley (2025), "
  "McKinnon (2017), Simmons (2014), and Zimman (2020) {} each of which nonetheless theorizes "
  "passing, reading, enregisterment, or appropriation. Because the set was selected for its "
  "topical proximity to the item, this absence is treated as evidence rather than as an "
  "artifact of sampling.".replace("{} ", f"{EN} ", 2))

# ============================================================
# 9. DISCUSSION
# ============================================================
H("9. Discussion")
P("Read together, the four literatures yield a coherent genealogy (Figure 4). A general "
  "slang verb of perception, attested from the late 1920s, was specialized within Black and "
  "Latinx ballroom and transgender communities into a high-stakes act of gender detection, "
  "the antonym of passing and realness. Theorized in transgender studies as a form of "
  "recognition with ambivalent stakes (Schiffer, 2022; Wong, 2025), the item then diffused "
  "outward through mediatized drag and social media, re-generalizing into a broad "
  "contemporary sense while undergoing the indexical erasure that attends the appropriation "
  "of AAVE and ballroom vernacular (Crowley, 2025; Ilbury & Walcott, 2026).")
FIGURE(f"{FIG}/fig4.png",
       "Figure 4. The item clock it at the intersection of four scholarly lenses, each of "
       "which engages it adjacently rather than as a primary object of study.", width=5.4)
P("The synthesis also makes a methodological point. That the field has theorized the "
  "politics of this item far more thoroughly than it has described the item itself is the "
  "review's central finding and its principal invitation to future work, summarized with "
  "confidence ratings in Table 2.")

TBL_CAP("Table 2. Principal claims, best supporting sources, source tier, and verification status.")
claim_rows = [
    ("Claim", "Best source(s)", "Tier", "Status"),
    ("Perception sense attested 1929–1942, predating ballroom use",
     "New Partridge; Berrey & Van den Bark (1942); Green's", "Lexicographic", "Secure"),
    ("Ballroom/trans use is a specialization, not the origin",
     "Inference from lexicographic chronology", "Inference", "Secure"),
    ("'Strike' sense (1941) is a separate branch from the perception sense",
     "OED", "Lexicographic", "Probable"),
    ("In trans studies, clocking is the antonym of passing",
     "Wong (2025); Schiffer (2022)", "Peer-reviewed", "Secure"),
    ("Passing-as-misrecognition / clocking-as-recognition inversion",
     "Wong (2025)", "Peer-reviewed", "Secure"),
    ("Diffusion best framed as racialized appropriation / erasure",
     "Ilbury (2020); Ilbury & Walcott (2026); Crowley (2025)", "Peer-reviewed", "Secure"),
    ("clock vs. reading vs. shade are distinct speech acts",
     "McKinnon (2017, 2023); Simmons (2014)", "Peer-reviewed", "Probable"),
    ("Regional origin (US vs. British) of perception sense",
     "OED vs. usage evidence", "Lexicographic", "Contested"),
    ("No dedicated study of clock; lemma absent in 6 of 7 full texts",
     "Lexical check (this review)", "Primary check", "Secure"),
]
t2 = doc.add_table(rows=0, cols=4); t2.style = "Table Grid"; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Inches(2.9), Inches(2.2), Inches(1.0), Inches(0.9)]
for i, r in enumerate(claim_rows):
    cells = t2.add_row().cells
    for c, txt, w in zip(cells, r, widths):
        cell_text(c, txt, bold=(i == 0), size=9)
        c.width = w
TBL_NOTE("Tiers: lexicographic = scholarly slang dictionaries; peer-reviewed = journal "
         "articles or scholarly monographs; inference = derived from other evidence; "
         "primary check = direct full-text examination conducted for this review.")

# ============================================================
# 10. LIMITATIONS
# ============================================================
H("10. Limitations")
P("This review is constrained by its subject's literature. No peer-reviewed study takes "
  "clock as its primary object, so the synthesis necessarily draws the item out of works "
  "centered elsewhere. The lexicographic datings rest on scholarly slang dictionaries "
  "consulted through catalogue and secondary evidence and warrant confirmation against the "
  "OED and Green's full entries, which would also resolve the open questions of regional "
  "origin and earliest citation. The full-text lexical check covered a purposive set of "
  "seven sources rather than an exhaustive corpus; a larger sample could refine, though it "
  "is unlikely to overturn, the lacuna it documents.")

# ============================================================
# 11. CONCLUSION
# ============================================================
H("11. Conclusion")
P("The verb to clock and its imperative clock it compress, in a single lexical item, three "
  "processes that sociolinguistics usually studies apart: documented semantic change, the "
  "enregisterment of a minoritized variety, and the racialized appropriation of that variety "
  "into a putatively neutral mainstream. The item's history is best told not as a story of "
  "ballroom invention but as one of specialization and re-generalization. That the field has "
  "theorized the politics of this item far more thoroughly than it has described the item "
  "itself is the review's central finding and its principal invitation to future work. To "
  "clock the tea, in the end, is to perform exactly the act of recognition this article "
  "recommends turning back upon the word itself.")

# ============================================================
# REFERENCES
# ============================================================
H("References")
refs = [
    "Agha, A. (2003). The social life of cultural value. Language & Communication, 23(3–4), 231–273. https://doi.org/10.1016/S0271-5309(03)00012-0",
    "Bailey, M. M. (2013). Butch queens up in pumps: Gender, performance, and ballroom culture in Detroit. University of Michigan Press.",
    "Barrett, R. (1998). Markedness and styleswitching in performances by African American drag queens. In C. Myers-Scotton (Ed.), Codes and consequences: Choosing linguistic varieties (pp. 139–161). Oxford University Press.",
    "Barrett, R. (2017). From drag queens to leathermen: Language, gender, and gay male subcultures. Oxford University Press.",
    "Beauchamp, T. (2019). Going stealth: Transgender politics and U.S. surveillance practices. Duke University Press.",
    "Berrey, L. V., & Van den Bark, M. (1942). The American thesaurus of slang. Thomas Y. Crowell.",
    f"Crowley, A. (2025). {LQ}And it just becomes queer slang{RQ}: Race, linguistic innovation, and appropriation within trans communities in the US South. Journal of Linguistic Anthropology, 35(2), e70008. https://doi.org/10.1111/jola.70008",
    "Dalzell, T., & Victor, T. (Eds.). (2006). The new Partridge dictionary of slang and unconventional English. Routledge.",
    f"Goldmark, M. (2015). National drag: The language of inclusion in RuPaul{AP}s Drag Race. GLQ: A Journal of Lesbian and Gay Studies, 21(4), 501–520. https://doi.org/10.1215/10642684-3123665",
    f"Green, J. (n.d.). Clock, v. In Green{AP}s dictionary of slang. Retrieved June 8, 2026, from https://greensdictofslang.com/entry/acsmoaa",
    "Greenhalgh, T., Thorne, S., & Malterud, K. (2018). Time to challenge the spurious hierarchy of systematic over narrative reviews? European Journal of Clinical Investigation, 48(6), e12931. https://doi.org/10.1111/eci.12931",
    f"Hodes, C., & Sandoval, J. (2018). RuPaul{AP}s Drag Race: A study in the commodification of white ruling-class femininity and the etiolation of drag. Studies in Costume & Performance, 3(2), 149–166. https://doi.org/10.1386/scp.3.2.149_1",
    "Homosaurus. (n.d.). Clocking (gender) [homoit0001781]. In Homosaurus (v3). Retrieved June 8, 2026, from https://homosaurus.org/v3/homoit0001781",
    f"Ilbury, C. (2020). {LQ}Sassy queens{RQ}: Stylistic orthographic variation in Twitter and the enregisterment of AAVE. Journal of Sociolinguistics, 24(2), 245–264. https://doi.org/10.1111/josl.12366",
    f"Ilbury, C., & Walcott, R. (2026). {LQ}Gen Z language? Y{AP}all mean AAVE{RQ}: The appropriation of African American Vernacular English as {LQ}TikTok language.{RQ} Journal of Sociolinguistics. Advance online publication. https://doi.org/10.1111/josl.70024",
    f"McKinnon, S. (2017). {LQ}Building a thick skin for each other{RQ}: The use of {LQ}reading{RQ} as an interactional practice of mock impoliteness in drag queen backstage talk. Journal of Language and Sexuality, 6(1), 90–127. https://doi.org/10.1075/jls.6.1.04mck",
    f"McKinnon, S. (2023). {LQ}Haute couture? More like haute glue!{RQ}: The discourse of the reading challenge in RuPaul{AP}s Drag Race. Journal of Language and Sexuality, 12(1), 73–97. https://doi.org/10.1075/jls.21013",
    "Oxford English Dictionary. (n.d.). Clock (n.1, v.). In Oxford English Dictionary. Retrieved June 8, 2026, from https://www.oed.com/dictionary/clock_n1",
    "Schiffer, D. J. (2022). Researching while trans: Being clocked and cooling cistress. Journal of Contemporary Ethnography, 51(5), 700–725. https://doi.org/10.1177/08912416221081870",
    "Serano, J. (2007). Whipping girl: A transsexual woman on sexism and the scapegoating of femininity. Seal Press.",
    f"Simmons, N. (2014). Speaking like a queen in RuPaul{AP}s Drag Race: Towards a speech code of American drag queens. Sexuality & Culture, 18(3), 630–648. https://doi.org/10.1007/s12119-013-9213-2",
    "Wong, J. (2025). Ambivalences of trans recognition. Hypatia, 40(2), 269–289. https://doi.org/10.1017/hyp.2024.61",
    "Zimman, L. (2020). Transgender language, transgender moment: Toward a trans linguistics. In K. Hall & R. Barrett (Eds.), The Oxford handbook of language and sexuality. Oxford University Press.",
]
for ref in refs:
    p = doc.add_paragraph(); p.alignment = J
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(4)
    _fmt(p.add_run(ref), size=10)

doc.save(OUT)
print("Saved:", OUT)
